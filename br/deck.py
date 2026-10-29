#!/usr/bin/env python3
"""Build the seven-slide template deck from a content JSON and figures.

Usage: deck.py --content deck/content.json --template deck/template.pptx --out deck/BeforeRoute.pptx

content.json keys: title, name, slides: {2..7: {title, kicker, bullets:[...], image:path|null,
image_box:[left,top,width,height] inches | null, table: {rows:[[...]]} (slide 5), caption}}
Bullets may be strings or {"text": str, "level": int, "bold": bool, "size": pt, "color": "RRGGBB"}.
Optional layout keys (all ignored when absent, so older content files build unchanged):
  bullets_uniform: true    every level-0 paragraph gets the template's first-paragraph bullet
  text_box: [l, t, w, h]   reposition the body placeholder (inches)
  images: [{"path": p, "box": [l, t, w]}, ...]   extra pictures (any slide, incl. 5)
  title_notes: str          speaker notes for slide 1;  notes_size: pt   font size for all speaker notes
  textboxes: [{"box": [l, t, w, h], "lines": [{"text", "size", "color", "bold", "align"}]}]  free text (any slide;
                             title_textboxes for slide 1)
  tiles: [{"box": [l, t, w, h], "lines": [...], "fill", "line"}]  native rounded tiles (big number + label)
  remove_body: true         drop the body placeholder (figure-only slides); slide 5: remove_caption, remove_table
  slide 5: caption_box [l, t, w, h] (word-wrapped caption), caption_size,
           table.col_widths [in, ...], table.row_height in, table.body_text_color "RRGGBB"
"""
from __future__ import annotations

import argparse
import copy
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def set_text(shape, lines, size=None, uniform_bullets=False, space=None):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]._p
    bullet_ppr = copy.deepcopy(p0.pPr) if (uniform_bullets and p0.pPr is not None) else None
    tf.clear()
    first = True
    for item in lines:
        if isinstance(item, str):
            item = {"text": item}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        lvl = int(item.get("level", 0))
        if bullet_ppr is not None and lvl == 0 and item.get("bullet", True):
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            p._p.insert(0, copy.deepcopy(bullet_ppr))
        p.level = lvl
        if space is not None:                        # optional tighter paragraph spacing (points before; none after)
            p.space_before = Pt(space)
            p.space_after = Pt(0)
        r = p.add_run()
        r.text = item["text"]
        if item.get("bold"):
            r.font.bold = True
        if item.get("color"):
            r.font.color.rgb = RGBColor.from_string(item["color"])
        sz = item.get("size", size)
        if sz:
            r.font.size = Pt(sz)


def add_images(slide, spec):
    for im in spec.get("images") or []:
        l, t, w = im["box"][:3]
        slide.shapes.add_picture(im["path"], Inches(l), Inches(t), width=Inches(w))


def add_textboxes(slide, spec):
    """Free text boxes (e.g. large callout numbers): {"box": [l, t, w, h], "lines": [{"text", "size", "color", "bold",
    "align": "l"|"c"|"r"}, ...]}; editable in PowerPoint like any text."""
    from pptx.enum.text import PP_ALIGN
    for tb in spec.get("textboxes") or []:
        l, t, w, h = tb["box"]
        shp = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = shp.text_frame
        tf.word_wrap = True
        for i, ln in enumerate(tb["lines"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}.get(ln.get("align", "l"), PP_ALIGN.LEFT)
            if ln.get("space") is not None:
                p.space_before = Pt(ln["space"])
            r = p.add_run()
            r.text = ln["text"]
            r.font.size = Pt(ln.get("size", 14))
            r.font.bold = bool(ln.get("bold"))
            if ln.get("color"):
                r.font.color.rgb = RGBColor.from_string(ln["color"])


def add_tiles(slide, spec):
    """Native rounded tiles (e.g. a big number over a short label): {"box": [l, t, w, h], "lines": [{"text", "size",
    "color", "bold", "align", "space"}], "fill": "RRGGBB", "line": "RRGGBB" | None, "anchor": "t"|"m"|"b"}."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    for tl in spec.get("tiles") or []:
        l, t, w, h = tl["box"]
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
        shp.adjustments[0] = tl.get("radius", 0.07)
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(tl.get("fill", "141414"))
        if tl.get("line"):
            shp.line.color.rgb = RGBColor.from_string(tl["line"])
            shp.line.width = Pt(tl.get("line_w", 1.0))
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(tl.get("pad", 0.12))
        tf.margin_top = tf.margin_bottom = Inches(tl.get("pad_v", 0.05))
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[tl.get("anchor", "m")]
        for i, ln in enumerate(tl["lines"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}.get(ln.get("align", tl.get("align", "l")),
                                                                         PP_ALIGN.LEFT)
            if ln.get("space") is not None:
                p.space_before = Pt(ln["space"])
            r = p.add_run()
            r.text = ln["text"]
            r.font.size = Pt(ln.get("size", 12))
            r.font.bold = bool(ln.get("bold"))
            r.font.color.rgb = RGBColor.from_string(ln.get("color", "FFFFFF"))


def remove_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def set_notes(slide, text, size=None):
    """Speaker notes; each newline starts a paragraph. size (pt) is optional: the default is the template's."""
    tf = slide.notes_slide.notes_text_frame
    tf.text = text
    if size:
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size)


def find(slide, name_prefix=None, ph_idx=None):
    for sh in slide.shapes:
        if ph_idx is not None and sh.is_placeholder and sh.placeholder_format.idx == ph_idx:
            return sh
        if name_prefix and sh.name.startswith(name_prefix):
            return sh
    return None


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--content", required=True)
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    c = json.load(open(a.content))
    prs = Presentation(a.template)
    slides = list(prs.slides)

    # slide 1
    s1 = slides[0]
    set_text(find(s1, ph_idx=0), [c["title"]])
    set_text(find(s1, ph_idx=1), [c.get("subtitle", ""), c["name"]])
    for key, idx in (("title_box", 0), ("subtitle_box", 1)):   # optional [left, top, width, height] in inches
        if c.get(key):
            sh = find(s1, ph_idx=idx)
            sh.left, sh.top, sh.width, sh.height = (Inches(v) for v in c[key])
    add_textboxes(s1, {"textboxes": c.get("title_textboxes")})  # optional callouts on the title slide
    nsz = c.get("notes_size")  # optional speaker-notes font size (pt) for every slide
    if c.get("title_notes"):  # optional speaker notes for the title slide
        set_notes(s1, c["title_notes"], nsz)

    for n in range(2, 8):
        spec = c["slides"].get(str(n))
        if not spec:
            continue
        s = slides[n - 1]
        if spec.get("title") is not None:
            set_text(find(s, ph_idx=0), [spec["title"]])
        if spec.get("kicker") is not None and find(s, ph_idx=10) is not None:
            set_text(find(s, ph_idx=10), [spec["kicker"]])
        body = find(s, ph_idx=1)
        if n == 5:
            # table placeholder is ph 1 on slide 5; chart is a separate shape
            tbl_shape = find(s, name_prefix="Table")
            chart_shape = find(s, name_prefix="Chart")
            cap = find(s, name_prefix="TextBox")
            if not spec.get("table") and tbl_shape is not None and spec.get("remove_table"):
                remove_shape(tbl_shape)
                tbl_shape = None
            if spec.get("table") and tbl_shape is not None:
                rows = spec["table"]["rows"]
                tbl = tbl_shape.table
                # grow/shrink rows: python-pptx has no add_row; clone the last row element
                while len(tbl.rows) < len(rows):
                    new_tr = copy.deepcopy(tbl._tbl.tr_lst[-1])
                    tbl._tbl.append(new_tr)
                while len(tbl.rows) > len(rows):
                    tbl._tbl.remove(tbl._tbl.tr_lst[-1])
                while len(tbl.columns) < len(rows[0]):
                    # clone last column (gridCol + each row's last cell)
                    grid = tbl._tbl.tblGrid
                    grid.append(copy.deepcopy(grid.gridCol_lst[-1]))
                    for tr in tbl._tbl.tr_lst:
                        tr.append(copy.deepcopy(tr.tc_lst[-1]))
                for j, w in enumerate(spec["table"].get("col_widths") or []):
                    if j < len(tbl.columns):
                        tbl.columns[j].width = Inches(w)
                if spec["table"].get("row_height"):
                    for row_ in tbl.rows:  # minimum height; rows still grow to fit their text
                        row_.height = Inches(spec["table"]["row_height"])
                body_rgb = spec["table"].get("body_text_color")
                for i, row in enumerate(rows):
                    for j, val in enumerate(row):
                        cell = tbl.cell(i, j)
                        cell.text = str(val)
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                r.font.size = Pt(spec["table"].get("font", 11))
                                if i == 0:
                                    r.font.bold = True
                                elif body_rgb:
                                    r.font.color.rgb = RGBColor.from_string(body_rgb)
                if spec["table"].get("box"):
                    l, t, w, h = spec["table"]["box"]
                    tbl_shape.left, tbl_shape.top, tbl_shape.width, tbl_shape.height = Inches(l), Inches(t), Inches(w), Inches(h)
            if chart_shape is not None:
                if spec.get("image"):
                    l, t, w, h = spec.get("image_box") or [7.1, 2.25, 5.3, 4.05]
                    remove_shape(chart_shape)
                    s.shapes.add_picture(spec["image"], Inches(l), Inches(t), width=Inches(w))
                else:
                    remove_shape(chart_shape)
            if cap is not None and spec.get("caption") is not None:
                if spec.get("caption_box"):
                    l, t, w, h = spec["caption_box"]
                    cap.left, cap.top, cap.width, cap.height = Inches(l), Inches(t), Inches(w), Inches(h)
                    cap.text_frame.word_wrap = True
                set_text(cap, spec["caption"] if isinstance(spec["caption"], list) else [spec["caption"]], size=spec.get("caption_size", 11))
            if cap is not None and spec.get("remove_caption"):
                remove_shape(cap)
            add_images(s, spec)
            add_tiles(s, spec)
            add_textboxes(s, spec)
            if spec.get("notes"):
                set_notes(s, spec["notes"], nsz)
            continue
        if body is not None and spec.get("remove_body"):     # a slide built only from figures and tiles
            remove_shape(body)
            body = None
        if body is not None:
            if spec.get("text_box"):
                l, t, w, h = spec["text_box"]
                body.left, body.top, body.width, body.height = Inches(l), Inches(t), Inches(w), Inches(h)
            if spec.get("image"):
                l, t, w, h = spec.get("image_box") or [7.0, 2.5, 5.7, 4.3]
                # shrink the text body to the left column
                body.width = Inches(spec.get("text_width", 6.2))
                s.shapes.add_picture(spec["image"], Inches(l), Inches(t), width=Inches(w))
            set_text(body, spec.get("bullets", []), size=spec.get("font"), uniform_bullets=bool(spec.get("bullets_uniform")),
                     space=spec.get("para_space"))
        add_images(s, spec)
        add_tiles(s, spec)
        add_textboxes(s, spec)
        if spec.get("notes"):
            set_notes(s, spec["notes"], nsz)
    cp = prs.core_properties  # the template carries someone else's title/author; set ours (they show in the PDF too)
    t = c.get("title", "BeforeRoute")
    cp.title = t.get("text", "BeforeRoute") if isinstance(t, dict) else str(t)
    cp.author = c.get("name", "")
    cp.last_modified_by = c.get("name", "")
    st = c.get("subtitle", "")
    cp.subject = (st.get("text", "") if isinstance(st, dict) else str(st))[:255]  # the core property's limit
    cp.comments = ""
    cp.keywords = ""
    prs.save(a.out)
    print(f"wrote {a.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
